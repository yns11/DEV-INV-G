"""Recalculer un classeur pour de vrai, et lire ce qu'il donne.

Un classeur de repli dont les formules sont fausses est pire qu'une absence de
repli : il s'ouvre, il affiche des chiffres, et rien ne dit qu'ils ne sont pas
ceux de l'application. Le seul contrôle qui vaille est donc celui-ci — faire
**recalculer** le fichier par un tableur, et comparer ce qu'il rend avec ce que
le moteur produit sur les mêmes données.

``xlsxwriter`` n'écrit pas de valeur en cache : les formules partent sans
résultat, et c'est le tableur qui les évalue à l'ouverture. LibreOffice, lui,
ne recalcule **pas** un fichier Excel par défaut — il fait confiance au cache,
qui est ici vide, et rend zéro partout. Il faut donc le lui demander, ce que
fait :func:`_profile` en posant ``OOXMLRecalcMode`` à « toujours » dans un
profil jetable.

LibreOffice n'est pas installé partout, et cette recette ne doit pas être une
condition pour lancer les contrôles. Les tests qui l'utilisent s'ignorent quand
il manque — comme ceux qui exigent un PostgreSQL — et les contrôles de
structure, eux, tournent toujours.
"""

from __future__ import annotations

import shutil
import subprocess
import tempfile
from pathlib import Path
from typing import Any

__all__ = ["AVAILABLE", "WHY_NOT", "recalculate"]

_SOFFICE = shutil.which("soffice") or shutil.which("libreoffice")

#: Le module Calc s'installe séparément du cœur de LibreOffice : ``soffice``
#: peut exister sans savoir ouvrir un tableur. Le vérifier ici évite un échec
#: illisible (« source file could not be loaded ») au milieu d'une recette.
_HAS_CALC = bool(_SOFFICE) and bool(
    list(Path("/usr/lib/libreoffice/program").glob("libsclo*"))
    or list(Path("/usr/lib/libreoffice/program").glob("libscfiltlo*"))
)

AVAILABLE = _HAS_CALC
WHY_NOT = (
    "LibreOffice Calc absent : le recalcul réel du classeur ne peut pas être "
    "vérifié ici (les contrôles de structure, eux, tournent)."
)

#: Le réglage qui force le recalcul à l'ouverture d'un fichier OOXML.
#: 0 = toujours, 1 = jamais (le défaut), 2 = demander.
_RECALC_ALWAYS = (
    '<item oor:path="/org.openoffice.Office.Calc/Formula/Load">'
    '<prop oor:name="OOXMLRecalcMode" oor:op="fuse">'
    "<value>0</value></prop></item>"
)

_TIMEOUT_S = 300


def recalculate(content: bytes) -> dict[str, list[list[Any]]]:
    """Faire recalculer *content* et rendre ``{feuille: lignes}``.

    Les valeurs rendues sont celles que le tableur a **calculées**, pas celles
    qu'on y a écrites : une formule fausse se voit ici et nulle part ailleurs.
    """
    import openpyxl

    with tempfile.TemporaryDirectory() as tmp:
        root = Path(tmp)
        source = root / "classeur.xlsx"
        source.write_bytes(content)
        out = root / "out"
        out.mkdir()
        profile = _profile(root)

        result = subprocess.run(
            [
                str(_SOFFICE), "--headless", "--norestore",
                f"-env:UserInstallation=file://{profile}",
                "--convert-to", "xlsx", "--outdir", str(out), str(source),
            ],
            capture_output=True, text=True, timeout=_TIMEOUT_S, check=False,
        )
        produced = out / "classeur.xlsx"
        if not produced.exists():
            raise RuntimeError(
                "Le tableur n'a pas rendu de fichier : "
                f"{result.stdout.strip()} {result.stderr.strip()}"
            )

        book = openpyxl.load_workbook(produced, data_only=True)
        return {
            name: [
                [cell.value for cell in row]
                for row in book[name].iter_rows()
            ]
            for name in book.sheetnames
        }


def _profile(root: Path) -> Path:
    """Un profil LibreOffice jetable, réglé pour recalculer.

    Créé par un premier démarrage à vide — le fichier de réglages n'existe pas
    avant — puis complété. Le faire à la main serait plus court et plus
    fragile : sa forme exacte change d'une version à l'autre.
    """
    profile = root / "profil"
    subprocess.run(
        [
            str(_SOFFICE), "--headless", "--norestore", "--terminate_after_init",
            f"-env:UserInstallation=file://{profile}",
        ],
        capture_output=True, text=True, timeout=_TIMEOUT_S, check=False,
    )
    registry = profile / "user" / "registrymodifications.xcu"
    if not registry.exists():
        raise RuntimeError("LibreOffice n'a pas créé de profil utilisateur.")
    text = registry.read_text()
    registry.write_text(text.replace("</oor:items>", _RECALC_ALWAYS + "</oor:items>"))
    return profile
