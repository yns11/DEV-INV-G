"""Les portefeuilles d'articles, et la bascule « uniquement mes références ».

Le besoin
---------
Une campagne porte quatre à cinq cents références et l'analyse se répartit
entre plusieurs personnes. L'application savait déjà répartir des
**emplacements** — c'est « mon périmètre », porté par les gestionnaires — et ne
savait pas répartir des **articles**. Ce n'est pas la même découpe : un acheteur
suit ses références partout où elles sont, quel que soit l'entrepôt qui les
range, et le périmètre d'un gestionnaire ne l'aide en rien. Sur la vue Écarts,
sur le stock ERP et sur l'écart backflush, la première chose à faire était donc
de retrouver les siennes dans la liste de tout le monde.

Ce que ces contrôles tiennent
-----------------------------
**Plusieurs personnes suivent la même référence.** L'identité fait partie de la
clé depuis la migration 034. Le premier choix — un propriétaire unique — avait
pour lui d'être sans ambiguïté, et décrivait mal l'organisation : un acheteur et
un contrôleur de gestion regardent les mêmes articles, sans que l'un soit le
propriétaire de l'autre. Les décomptes changent avec la clé, et c'est là que ce
genre de changement laisse des traces : la somme des « X références » par
personne dépasse le nombre de références, tandis que « sans propriétaire » doit
continuer à compter des références **distinctes**. L'erreur inverse annoncerait
moins d'orphelines qu'il n'y en a, c'est-à-dire rassurerait au lieu d'alerter.

**Le tableau se charge comme les autres grilles**, avec deux portées. Entre les
références, une fusion : un fichier de trente références ne dit rien des quatre
cent cinquante autres, et les effacer ferait d'une correction ciblée une remise
à zéro. Sur une référence citée, un remplacement : le fichier fait foi pour
elle, ce qui est la seule façon de retirer *une* personne d'une référence
partagée. Une adresse vide est le cas limite de cette règle.

**Le filtre est résolu côté serveur**, à partir de l'identité que la plateforme
transmet — comme « mon périmètre ». Le navigateur n'envoie qu'un booléen, ce qui
est la seule forme où demander le portefeuille d'un autre est impossible plutôt
qu'interdit.

**Le filtre s'applique avant les totaux, avant les tris et avant les plafonds.**
C'est là que ce genre de fonctionnalité se trompe sans qu'on le voie : « les
vingt plus gros écarts, dont ceux qui sont à moi » rend parfois zéro ligne sur
une liste où l'on en attend vingt, et un total calculé sur des lignes que
l'écran ne montre pas est le genre de chiffre qu'on passe une matinée à ne pas
retrouver.

Ce que ce n'est pas
-------------------
**Une habilitation.** Un portefeuille filtre l'affichage et n'interdit rien :
chacun garde le droit d'agir sur toutes les références. Le droit d'écrire vient
d'être déclaré gestionnaire, et de rien d'autre.
"""

from __future__ import annotations

import datetime as dt
import uuid
from decimal import Decimal

import pytest
from tests.early_count_db import disposable_database, make_campaign

from inventory.domain.enums import CampaignStatus
from inventory.domain.models import (
    BackflushLine,
    BookStockLine,
    Campaign,
    Item,
    ItemPortfolio,
)
from inventory.errors import FrozenError

pytestmark = pytest.mark.postgres

ANNE = "anne.durand@usine.fr"
BORIS = "boris.leroy@usine.fr"

WH, LOC = "B06", "A-01"

#: Trois références, dont deux à Anne. La troisième est ce qui fait la
#: différence entre « le filtre marche » et « le filtre laisse tout passer ».
A_ANNE = ("P-100", "P-200")
A_BORIS = ("P-300",)
TOUTES = A_ANNE + A_BORIS


@pytest.fixture(scope="module")
def db():
    with disposable_database("inventaire_portefeuilles") as database:
        yield database


@pytest.fixture
def campaign(db):
    campaign_id = make_campaign(db, f"PF-{uuid.uuid4().hex[:8]}")
    with db.transaction() as conn:
        conn.execute(
            "UPDATE campaign SET status='ANALYSIS' WHERE id=%s", (campaign_id,)
        )
    return _campaign(campaign_id, CampaignStatus.ANALYSIS)


def _campaign(campaign_id: str, status: CampaignStatus) -> Campaign:
    return Campaign(
        id=campaign_id,
        code=f"PF-{campaign_id[:8]}",
        label="",
        count_date=dt.date(2026, 9, 12),
        status=status,
        created_by="test",
        created_at=dt.datetime(2026, 9, 1, tzinfo=dt.UTC),
    )


def _ctx(db, actor: str):
    from inventory.config import get_settings
    from inventory.services.context import ServiceContext

    return ServiceContext(actor=actor, db=db, settings=get_settings())


@pytest.fixture
def ctx(db, campaign, monkeypatch):
    """Le contexte d'Anne, et les trois articles de la campagne.

    Les prix sont distincts et croissants : un écart valorisé qui tomberait
    juste par hasard ne prouverait rien du tri.

    L'archive est forcée sur la base : le chargement d'un portefeuille dépose sa
    pièce justificative comme tout import, et le dos « volume Unity Catalog »
    demande un espace de travail Databricks que ces contrôles n'ont pas.
    """
    from inventory.config import Settings, get_settings

    monkeypatch.setenv("INV_EVIDENCE_STORE", "lakebase")
    get_settings.cache_clear()
    monkeypatch.setattr(
        Settings, "evidence_configured", property(lambda self: True), raising=False
    )
    context = _ctx(db, ANNE)
    context.referentials.upsert_items(
        [
            Item(
                campaign_id=campaign.id,
                item_number=number,
                name=f"Article {number}",
                std_price=Decimal(10 * (index + 1)),
            )
            for index, number in enumerate(TOUTES)
        ],
        actor="test",
    )
    return context


@pytest.fixture
def client(db):
    """L'application réelle, sur la base jetable du module.

    ``disposable_database`` déplace ``PGDATABASE`` pendant toute la durée du
    module : l'application construite ici parle donc à la même base que les
    fixtures, et ce qu'elles sèment est visible depuis les routes.

    Le point d'ancrage est remis à zéro en sortant. Sans cela, les connexions
    ouvertes par l'application survivent au client et la base jetable refuse de
    se laisser supprimer — « is being accessed by other users », à la fin d'un
    module dont tous les contrôles sont pourtant passés.
    """
    from fastapi.testclient import TestClient

    from inventory.api import create_app
    from inventory.db.engine import reset_database

    with TestClient(create_app()) as running:
        yield running
    reset_database()


@pytest.fixture
def seeded(db, campaign, ctx):
    """Une campagne complète : trois articles, leur stock ERP, leurs propriétaires."""
    _stock(ctx, campaign)
    _attribuer(ctx, campaign, ("P-100", ANNE), ("P-200", ANNE), ("P-300", BORIS))
    return campaign.id


@pytest.fixture
def chef(db, ctx):
    """Le contexte du propriétaire, pour tout ce qui écrit.

    Anne n'est ni propriétaire ni gestionnaire, et c'est **voulu** : un
    portefeuille n'est pas une habilitation, et quelqu'un qui suit des
    références n'a pas pour autant le droit d'écrire. La séparation des deux
    contextes fait donc partie de ce qui est vérifié ici — Anne lit le sien sans
    pouvoir toucher au tableau.
    """
    return _ctx(db, "test")


@pytest.fixture
def service(chef):
    from inventory.services import ImportService

    return ImportService(chef)


def _attribuer(ctx, campaign: Campaign, *attributions: tuple[str, str]) -> None:
    """Poser des couples (référence, personne).

    Des couples et non un dictionnaire : depuis le partage, une référence en
    porte plusieurs, et un dictionnaire indexé par référence ne saurait pas
    l'écrire — le contrôle serait alors incapable d'exprimer ce qu'il vérifie.
    """
    ctx.portfolios.upsert(
        campaign.id,
        [
            ItemPortfolio(campaign_id=campaign.id, item_number=number, actor=actor)
            for number, actor in attributions
        ],
        actor="test",
    )


def _stock(ctx, campaign: Campaign) -> None:
    """Une ligne de stock ERP par référence, de valeur croissante."""
    ctx.book_stock.replace(
        campaign.id,
        [
            BookStockLine(
                campaign_id=campaign.id,
                item_number=number,
                warehouse_id=WH,
                location_id=LOC,
                qty=Decimal(10),
                unit_cost=Decimal(10 * (index + 1)),
                reference_date=dt.date(2026, 9, 12),
            )
            for index, number in enumerate(TOUTES)
        ],
        batch_id=None,
    )


def _backflush(ctx, campaign: Campaign) -> None:
    ctx.backflush.replace(
        campaign.id,
        [
            BackflushLine(
                campaign_id=campaign.id,
                item_number=number,
                period_start=dt.date(2026, 8, 31),
                period_end=dt.date(2026, 9, 7),
                net_qty=Decimal(index + 1),
                theoretical_qty=Decimal(100),
                actual_qty=Decimal(100 - index - 1),
                week_count=1,
            )
            for index, number in enumerate(TOUTES)
        ],
        batch_id=None,
    )


# --------------------------------------------------------------------------- #
# 1. Écrire le tableau
# --------------------------------------------------------------------------- #


def _couples(ctx, campaign: Campaign) -> set[tuple[str, str]]:
    return {(r.item_number, r.actor) for r in ctx.portfolios.list(campaign.id)}


class TestUneReferenceSeSuitAPlusieurs:
    """La forme retenue : l'identité fait partie de la clé.

    Le premier choix était un propriétaire unique — sans ambiguïté, et décrivant
    mal l'organisation. Ces contrôles tiennent le partage lui-même ; ceux de la
    classe suivante tiennent la façon dont un fichier l'écrit.
    """

    def test_deux_personnes_suivent_la_meme(self, ctx, campaign):
        _attribuer(ctx, campaign, ("P-100", ANNE), ("P-100", BORIS))

        assert _couples(ctx, campaign) == {("P-100", ANNE), ("P-100", BORIS)}

    def test_et_chacune_la_voit_dans_ses_références(self, db, ctx, campaign):
        """Ce à quoi le partage sert : la bascule la rend aux deux.

        Sans cela la clé élargie ne serait qu'une ligne de plus en base — le
        filtre, lui, continuerait de n'en désigner qu'une.
        """
        _attribuer(ctx, campaign, ("P-100", ANNE), ("P-100", BORIS))

        assert ctx.portfolios.items_of(campaign.id, ANNE) == frozenset({"P-100"})
        assert ctx.portfolios.items_of(campaign.id, BORIS) == frozenset({"P-100"})

    def test_la_même_paire_deux_fois_ne_fait_qu_une_ligne(self, ctx, campaign):
        """Un tableur répète volontiers la même paire, et deux fois « P-100 est
        à Anne » dit une seule chose."""
        _attribuer(ctx, campaign, ("P-100", ANNE), ("P-100", ANNE))

        assert _couples(ctx, campaign) == {("P-100", ANNE)}

    def test_les_décomptes_par_personne_se_cumulent(self, ctx, campaign):
        """Et la somme dépasse le nombre de références. C'est exact : chacun en
        suit bien autant."""
        _attribuer(ctx, campaign, ("P-100", ANNE), ("P-100", BORIS), ("P-200", ANNE))

        par_personne = {
            row["actor"]: row["items"]
            for row in ctx.portfolios.counts_by_actor(campaign.id)
        }
        assert par_personne == {ANNE: 2, BORIS: 1}

    def test_mais_les_références_couvertes_restent_distinctes(self, ctx, campaign):
        """Le décompte qui ne doit surtout pas suivre le précédent.

        Compter les lignes ferait passer une référence suivie à deux pour deux
        références couvertes, et l'écran annoncerait moins d'orphelines qu'il
        n'y en a — l'erreur qui rassure au lieu d'alerter.
        """
        _attribuer(ctx, campaign, ("P-100", ANNE), ("P-100", BORIS), ("P-200", ANNE))

        assert ctx.portfolios.assigned_items(campaign.id) == 2

    def test_les_propriétaires_d_une_référence_se_suivent(self, ctx, campaign):
        """La grille ouvre sur l'ordre qui montre le partage.

        Par référence d'abord : les deux personnes d'une référence partagée
        arrivent côte à côte. Ordonner par personne les séparerait de plusieurs
        écrans, et donnerait à lire deux portefeuilles indépendants là où il y a
        un partage — c'est-à-dire cacherait précisément ce que cette version
        ajoute.
        """
        _attribuer(ctx, campaign, ("P-200", ANNE), ("P-100", BORIS), ("P-100", ANNE))

        rows = ctx.portfolios.list(campaign.id)
        assert [(r.item_number, r.actor) for r in rows] == [
            ("P-100", ANNE), ("P-100", BORIS), ("P-200", ANNE),
        ]

    def test_l_écran_le_dit_de_la_même_façon(self, ctx, campaign):
        """Trois références connues, deux couvertes : une orpheline."""
        from inventory.services.portfolio_service import PortfolioService

        _attribuer(ctx, campaign, ("P-100", ANNE), ("P-100", BORIS), ("P-200", ANNE))

        vue = PortfolioService(ctx).overview(campaign)
        assert len(vue["rows"]) == 3
        assert vue["items"] == 2
        assert vue["unassigned"] == 1


class TestLeTableauSeChargeEtFusionne:
    def test_une_attribution_se_pose(self, ctx, campaign):
        _attribuer(ctx, campaign, ("P-100", ANNE))

        rows = ctx.portfolios.list(campaign.id)
        assert [(r.item_number, r.actor) for r in rows] == [("P-100", ANNE)]

    def test_un_second_chargement_ne_efface_pas_le_premier(self, ctx, campaign):
        """Le cas qui décide de la portée : une fusion **entre** les références.

        Un fichier de trente références ne dit rien des quatre cent cinquante
        autres. Les effacer parce qu'il ne les mentionne pas ferait d'une
        correction ciblée une remise à zéro — et sur une table tenue dans un
        tableur, on charge par morceaux.
        """
        _attribuer(ctx, campaign, ("P-100", ANNE), ("P-200", ANNE))
        _attribuer(ctx, campaign, ("P-300", BORIS))

        assert {r.item_number for r in ctx.portfolios.list(campaign.id)} == set(TOUTES)

    def test_une_référence_citée_est_remplacée(self, ctx, campaign):
        """Et l'autre portée : sur une référence citée, le fichier fait foi.

        C'est ce qui permet de retirer *une* personne d'une référence partagée —
        on la recharge avec la liste voulue. Une fusion pure ne saurait
        qu'ajouter, et obligerait à tout vider pour enlever quelqu'un.
        """
        _attribuer(ctx, campaign, ("P-100", ANNE), ("P-100", BORIS))
        _attribuer(ctx, campaign, ("P-100", ANNE))

        assert _couples(ctx, campaign) == {("P-100", ANNE)}

    def test_une_reference_change_de_main(self, ctx, campaign):
        _attribuer(ctx, campaign, ("P-100", ANNE))
        _attribuer(ctx, campaign, ("P-100", BORIS))

        assert _couples(ctx, campaign) == {("P-100", BORIS)}

    def test_une_adresse_vide_retire_toutes_les_attributions(self, ctx, campaign):
        """Le cas limite du remplacement, et non une règle à part : une
        référence citée sans personne n'est à personne. Sur une référence
        partagée, elle les retire donc tous."""
        _attribuer(ctx, campaign, ("P-100", ANNE), ("P-100", BORIS), ("P-200", ANNE))
        _attribuer(ctx, campaign, ("P-100", ""))

        assert _couples(ctx, campaign) == {("P-200", ANNE)}

    def test_l_adresse_ecrite_est_rangee_en_minuscules(self, ctx, campaign):
        """Un annuaire qui écrit « Prenom.Nom@ » un jour et « prenom.nom@ » le
        lendemain ne doit pas produire deux portefeuilles — dont l'un des deux
        ne rendrait jamais rien."""
        _attribuer(ctx, campaign, ("P-100", "  Anne.DURAND@usine.FR  "))

        assert [r.actor for r in ctx.portfolios.list(campaign.id)] == [ANNE]

    def test_et_l_adresse_lue_aussi(self, db, ctx, campaign):
        """Les deux bouts, parce que ce sont deux sources différentes.

        L'écriture vient d'un tableur, la lecture d'un en-tête que la plateforme
        transmet à chaque requête — et rien ne garantit que les deux emploient la
        même casse. Le jour où la connexion rend « Anne.Durand@ », le tableau est
        rangé en minuscules et le filtre d'Anne ne trouve plus rien : un écran
        vide, sans erreur, qu'on met une demi-journée à comprendre.
        """
        from inventory.services.portfolio_service import PortfolioService

        _attribuer(ctx, campaign, ("P-100", ANNE))

        connectee = PortfolioService(_ctx(db, " Anne.DURAND@usine.FR "))
        assert connectee.mine(campaign) == frozenset({"P-100"})

    def test_tout_retirer_est_un_geste_a_part(self, ctx, chef, campaign):
        from inventory.services.portfolio_service import clear_portfolios

        _attribuer(ctx, campaign, *((n, ANNE) for n in TOUTES))
        removed = clear_portfolios(chef, campaign)

        assert removed == 3
        assert ctx.portfolios.list(campaign.id) == []


class TestLeChargementPasseParLePipelineDesGrilles:
    """Fichier, collage ou saisie — le même contrat, comme toutes les grilles."""

    LIGNES = "Article\tE-mail\nP-100\tanne.durand@usine.fr\nP-200\tANNE.DURAND@usine.fr\n"

    def test_un_collage_charge_le_tableau(self, service, ctx, campaign):
        outcome = service.import_portfolios(campaign, mode="paste", text=self.LIGNES)

        assert outcome.rows_accepted == 2
        assert ctx.portfolios.items_of(campaign.id, ANNE) == frozenset(A_ANNE)

    def test_le_contrat_est_declare_comme_les_autres(self):
        from inventory.ingest.contracts import CONTRACTS

        assert "portfolios" in CONTRACTS

    def test_la_clé_naturelle_porte_aussi_l_identité(self):
        """Sans elle, les deux lignes d'une référence partagée seraient
        signalées en doublon — et le chargement qui la partage se ferait donc
        toujours sous un avertissement, ce qu'on apprend à ignorer."""
        from inventory.ingest.contracts import CONTRACTS

        assert CONTRACTS["portfolios"].natural_key == ("item_number", "actor")

    def test_deux_lignes_pour_une_référence_la_partagent(self, service, ctx, campaign):
        """La forme sous laquelle le partage arrive réellement : deux lignes du
        même fichier, et non deux chargements."""
        outcome = service.import_portfolios(
            campaign, mode="paste",
            text=f"Article\tE-mail\nP-100\t{ANNE}\nP-100\t{BORIS}\n",
        )

        assert outcome.rows_accepted == 2
        assert not outcome.warnings, "un partage n'est pas un doublon"
        assert _couples(ctx, campaign) == {("P-100", ANNE), ("P-100", BORIS)}

    def test_le_rapport_compte_les_références_et_les_partages(self, service, campaign):
        """Deux lignes sur une référence, une sur une autre : trois attributions,
        deux références, un partage. Le premier chiffre qu'on vérifie après un
        chargement est « ai-je bien touché mes références ? », et la hauteur du
        fichier n'y répond plus."""
        outcome = service.import_portfolios(
            campaign, mode="paste",
            text=f"Article\tE-mail\nP-100\t{ANNE}\nP-100\t{BORIS}\nP-200\t{ANNE}\n",
        )

        assert outcome.rows_accepted == 3
        assert outcome.details["items"] == 2
        assert outcome.details["shared"] == 1

    def test_un_rechargement_retire_quelqu_un_d_une_référence_partagée(
        self, service, ctx, campaign
    ):
        """Le geste que la forme « remplacement par référence » rend possible,
        et qu'une fusion pure interdirait : Boris s'en va, Anne reste, et les
        autres références ne bougent pas."""
        service.import_portfolios(
            campaign, mode="paste",
            text=f"Article\tE-mail\nP-100\t{ANNE}\nP-100\t{BORIS}\nP-200\t{BORIS}\n",
        )
        service.import_portfolios(
            campaign, mode="paste", text=f"Article\tE-mail\nP-100\t{ANNE}\n",
        )

        assert _couples(ctx, campaign) == {("P-100", ANNE), ("P-200", BORIS)}

    def test_une_reference_hors_referentiel_est_acceptee_et_signalee(
        self, service, ctx, campaign
    ):
        """Le tableau se prépare souvent avant que les articles ne soient
        chargés. La refuser bloquerait un chargement légitime ; la passer sous
        silence laisserait une coquille invisible jusqu'au jour où le filtre ne
        rend rien."""
        outcome = service.import_portfolios(
            campaign, mode="paste",
            text="Article\tE-mail\nP-100\tanne.durand@usine.fr\nP-999\tanne.durand@usine.fr\n",
        )

        assert outcome.details["unknownItems"] == 1
        assert outcome.details["unknownItemNumbers"] == ["P-999"]
        assert ctx.portfolios.items_of(campaign.id, ANNE) == frozenset({"P-100", "P-999"})

    def test_le_rapport_dit_combien_de_personnes(self, service, campaign):
        outcome = service.import_portfolios(
            campaign, mode="paste",
            text=(
                "Article\tE-mail\n"
                f"P-100\t{ANNE}\nP-200\t{BORIS}\nP-300\t\n"
            ),
        )

        assert outcome.details["actors"] == 2
        assert outcome.details["cleared"] == 1

    def test_le_chargement_est_rejouable(self):
        """Un chargement de portefeuilles se rejoue comme les autres.

        La table des cibles est la même : un import qui archive son fichier
        sans figurer ici serait rejouable par l'écran et refusé par le serveur.
        """
        from inventory.services.import_replay import TARGET_METHODS

        assert TARGET_METHODS["portfolios"] == "import_portfolios"


class TestLaGardeEstCelleDesGestionnaires:
    """Même nature de décision que les affectations — qui s'occupe de quoi — et
    elle bouge aux mêmes moments : jusqu'à la clôture comprise."""

    @pytest.mark.parametrize(
        "status",
        [CampaignStatus.PREPARATION, CampaignStatus.COUNTING, CampaignStatus.ANALYSIS],
    )
    def test_elle_reste_ouverte_jusqu_a_la_cloture(self, service, campaign, status):
        outcome = service.import_portfolios(
            _campaign(campaign.id, status), mode="paste",
            text=f"Article\tE-mail\nP-100\t{ANNE}\n",
        )

        assert outcome.rows_accepted == 1

    def test_elle_se_ferme_a_la_cloture(self, service, campaign):
        with pytest.raises(FrozenError):
            service.import_portfolios(
                _campaign(campaign.id, CampaignStatus.CLOSED), mode="paste",
                text=f"Article\tE-mail\nP-100\t{ANNE}\n",
            )

    def test_et_vider_aussi(self, chef, campaign):
        from inventory.services.portfolio_service import clear_portfolios

        with pytest.raises(FrozenError):
            clear_portfolios(chef, _campaign(campaign.id, CampaignStatus.CLOSED))


# --------------------------------------------------------------------------- #
# 2. Le relire comme un filtre
# --------------------------------------------------------------------------- #


class TestLeFiltreEstResoluCoteServeur:
    def test_chacun_lit_le_sien(self, db, ctx, campaign):
        """La même requête, deux réponses — parce que l'identité vient du
        contexte et non d'un paramètre. C'est ce qui rend impossible, et pas
        seulement interdit, de demander le portefeuille d'un autre."""
        from inventory.services.portfolio_service import PortfolioService

        _attribuer(ctx, campaign, ("P-100", ANNE), ("P-200", ANNE), ("P-300", BORIS))

        assert PortfolioService(ctx).mine(campaign) == frozenset(A_ANNE)
        assert PortfolioService(_ctx(db, BORIS)).mine(campaign) == frozenset(A_BORIS)

    def test_sans_bascule_rien_n_est_filtre(self, ctx, campaign):
        from inventory.services.portfolio_service import portfolio_filter

        _attribuer(ctx, campaign, ("P-100", ANNE))

        assert portfolio_filter(ctx, campaign, mine=False) is None

    def test_un_portefeuille_vide_ne_rend_pas_tout(self, ctx, campaign):
        """``None`` et l'ensemble vide disent deux choses différentes, et les
        confondre ferait mentir l'écran : sans filtre on voit tout, avec un
        filtre et aucun portefeuille on ne voit **rien**, et l'interface doit
        pouvoir le dire plutôt que de rendre la liste entière."""
        from inventory.services.portfolio_service import portfolio_filter

        assert portfolio_filter(ctx, campaign, mine=True) == frozenset()


class TestLeStockErpSuitLeFiltre:
    def test_il_ne_garde_que_les_siennes(self, ctx, campaign):
        from inventory.services import ReferentialService

        _stock(ctx, campaign)
        _attribuer(ctx, campaign, ("P-100", ANNE), ("P-200", ANNE), ("P-300", BORIS))

        view = ReferentialService(ctx).book_stock(
            campaign, only_items=frozenset(A_ANNE)
        )

        assert {line.item_number for line in view.lines} == set(A_ANNE)

    def test_le_total_porte_sur_ce_que_l_ecran_montre(self, ctx, campaign):
        """Un total calculé sur des lignes que la grille ne montre pas est le
        genre de chiffre qu'on passe une matinée à ne pas retrouver."""
        from inventory.services import ReferentialService

        _stock(ctx, campaign)

        entier = ReferentialService(ctx).book_stock(campaign)
        filtre = ReferentialService(ctx).book_stock(
            campaign, only_items=frozenset(A_ANNE)
        )

        assert entier.total_value == pytest.approx(600.0)  # 10×(10+20+30)
        assert filtre.total_value == pytest.approx(300.0)  # 10×(10+20)

    def test_le_top_se_calcule_apres_le_filtre(self, ctx, campaign):
        """« Les plus grosses lignes, dont celles qui sont à moi » rendrait
        parfois zéro ligne sur une liste où l'on en attend une."""
        from inventory.services import ReferentialService

        _stock(ctx, campaign)

        view = ReferentialService(ctx).book_stock(
            campaign, top=1, only_items=frozenset(A_ANNE)
        )

        assert [line.item_number for line in view.lines] == ["P-200"]
        assert view.top_share == pytest.approx(200.0 / 300.0)


class TestLesEcartsSuiventLeFiltre:
    def test_ils_ne_gardent_que_les_siennes(self, ctx, campaign):
        from inventory.services import AnalysisService

        _stock(ctx, campaign)

        rows = AnalysisService(ctx).top_variances(
            campaign, only_items=frozenset(A_ANNE)
        )

        assert {row["itemNumber"] for row in rows} == set(A_ANNE)

    def test_le_plafond_s_applique_apres_le_filtre(self, ctx, campaign):
        """Sans cela, « les vingt plus gros écarts, dont les miens » rend
        parfois aucune ligne sur un écran qui en attend vingt."""
        from inventory.services import AnalysisService

        _stock(ctx, campaign)

        rows = AnalysisService(ctx).top_variances(
            campaign, limit=1, only_items=frozenset(A_ANNE)
        )

        assert len(rows) == 1
        assert rows[0]["itemNumber"] in A_ANNE


class TestLEcartBackflushSuitLeFiltre:
    def test_les_lignes_sont_filtrees(self, ctx, campaign):
        from inventory.services import AnalysisService

        _stock(ctx, campaign)
        _backflush(ctx, campaign)

        view = AnalysisService(ctx).backflush(campaign, only_items=frozenset(A_ANNE))

        assert {row["itemNumber"] for row in view["rows"]} == set(A_ANNE)

    def test_le_bandeau_suit_la_grille(self, ctx, campaign):
        """Les trois premières cartes forment une soustraction — écart moins
        part expliquée égale inexpliqué. Des totaux de campagne au-dessus d'une
        grille filtrée ne tomberaient plus, et c'est l'écran où cela se verrait
        le plus."""
        from inventory.services import AnalysisService

        _stock(ctx, campaign)
        _backflush(ctx, campaign)

        entier = AnalysisService(ctx).backflush(campaign)
        filtre = AnalysisService(ctx).backflush(campaign, only_items=frozenset(A_ANNE))

        assert entier["kpis"]["backflushLineCount"] == 3
        assert filtre["kpis"]["backflushLineCount"] == 2


class TestLExportEmporteLaBascule:
    """Un fichier qui ne contiendrait pas ce qu'on avait sous les yeux au moment
    de cliquer serait le genre d'écart qu'on ne découvre qu'en réunion."""

    def test_le_classeur_ne_porte_que_les_siennes(self, ctx, campaign):
        from openpyxl import load_workbook

        from inventory.services import ReportService

        _stock(ctx, campaign)

        payload, _ = ReportService(ctx).variance_export(
            campaign, only_items=frozenset(A_ANNE)
        )
        import io

        workbook = load_workbook(io.BytesIO(payload))
        texte = "\n".join(
            str(cell.value)
            for sheet in workbook.worksheets
            for row in sheet.iter_rows()
            for cell in row
            if cell.value is not None
        )

        assert "P-100" in texte
        assert "P-300" not in texte

    def test_et_la_provenance_dit_sur_quel_perimetre(self, ctx, campaign):
        """Un total sur un portefeuille et un total sur la campagne portent le
        même titre et ne disent pas la même chose."""
        from openpyxl import load_workbook

        from inventory.services import ReportService

        _stock(ctx, campaign)

        payload, _ = ReportService(ctx).variance_export(
            campaign, only_items=frozenset(A_ANNE)
        )
        import io

        workbook = load_workbook(io.BytesIO(payload))
        texte = "\n".join(
            str(cell.value)
            for sheet in workbook.worksheets
            for row in sheet.iter_rows()
            for cell in row
            if cell.value is not None
        )

        assert "mon portefeuille" in texte


# --------------------------------------------------------------------------- #
# 3. Le câblage des routes
# --------------------------------------------------------------------------- #


class TestLesRoutesSontBranchees:
    """La classe de défaut de ce dépôt : un service écrit et que rien n'atteint."""

    @staticmethod
    def _routes():
        from inventory.api.app import create_app

        return {
            (route.path, method): route
            for route in create_app().routes
            if getattr(route, "methods", None)
            for method in route.methods
        }

    def test_le_tableau_se_lit_et_se_vide(self):
        chemin = "/api/campaigns/{campaign_id}/portfolios"
        routes = self._routes()

        assert (chemin, "GET") in routes
        assert (chemin, "DELETE") in routes

    @pytest.mark.parametrize(
        "chemin",
        [
            "/api/campaigns/{campaign_id}/book-stock",
            "/api/campaigns/{campaign_id}/analysis/variances",
            "/api/campaigns/{campaign_id}/analysis/backflush",
            "/api/campaigns/{campaign_id}/reports/variances.xlsx",
            "/api/campaigns/{campaign_id}/reports/variances.pdf",
        ],
    )
    def test_les_cinq_lectures_declarent_la_bascule(self, chemin):
        """Trois grilles et les deux exports qu'elles offrent.

        La forme seulement : qu'elle *fasse* quelque chose est vérifié plus bas,
        en traversant la route — déclarer le paramètre et l'ignorer donne
        exactement l'écran qui montre tout sous une bascule allumée.
        """
        route = self._routes()[(chemin, "GET")]
        noms = {p.name for p in route.dependant.query_params}

        assert "mine" in noms

    @pytest.mark.parametrize(
        "chemin",
        [
            "/api/campaigns/{campaign_id}/book-stock",
            "/api/campaigns/{campaign_id}/analysis/variances",
            "/api/campaigns/{campaign_id}/analysis/backflush",
            "/api/campaigns/{campaign_id}/reports/variances.xlsx",
            "/api/campaigns/{campaign_id}/reports/variances.pdf",
        ],
    )
    def test_et_aucune_ne_laisse_nommer_quelqu_un(self, chemin):
        """Le navigateur n'envoie qu'un booléen. Un paramètre qui porterait une
        adresse ferait du filtre une porte : il suffirait de la changer pour
        lire le portefeuille d'un autre."""
        route = self._routes()[(chemin, "GET")]
        noms = {p.name for p in route.dependant.query_params}

        assert not (noms & {"actor", "email", "user", "utilisateur"})


class TestLaBasculeTraverseReellementLaRoute:
    """Le paramètre est déclaré — mais est-il branché ?

    Une route peut accepter ``mine`` et l'ignorer : le contrat est respecté, le
    client généré compile, l'écran allume sa bascule et montre tout. C'est la
    classe de défaut de ce dépôt, et la seule façon de l'exclure est de faire la
    requête pour de bon, avec l'en-tête d'identité que la plateforme transmet.
    """

    @staticmethod
    def _get(client, url: str, *, actor: str):
        return client.get(url, headers={"X-Forwarded-Email": actor})

    def test_le_stock_erp_ne_rend_que_les_siennes(self, client, seeded):
        entier = self._get(
            client, f"/api/campaigns/{seeded}/book-stock", actor=ANNE
        ).json()
        filtre = self._get(
            client, f"/api/campaigns/{seeded}/book-stock?mine=true", actor=ANNE
        ).json()

        assert entier["total"] == 3
        assert {row["item_number"] for row in filtre["rows"]} == set(A_ANNE)

    def test_les_ecarts_ne_rendent_que_les_siennes(self, client, seeded):
        rows = self._get(
            client,
            f"/api/campaigns/{seeded}/analysis/variances?mine=true",
            actor=ANNE,
        ).json()

        assert {row["itemNumber"] for row in rows} == set(A_ANNE)

    def test_et_chacun_obtient_les_siennes_sur_la_meme_adresse(self, client, seeded):
        """La même URL, deux réponses : l'identité vient de l'en-tête.

        C'est ce qui rend *impossible*, et pas seulement interdit, de lire le
        portefeuille d'un autre — il n'y a rien à changer dans l'adresse pour
        l'obtenir.
        """
        url = f"/api/campaigns/{seeded}/analysis/variances?mine=true"

        anne = self._get(client, url, actor=ANNE).json()
        boris = self._get(client, url, actor=BORIS).json()

        assert {row["itemNumber"] for row in anne} == set(A_ANNE)
        assert {row["itemNumber"] for row in boris} == set(A_BORIS)
